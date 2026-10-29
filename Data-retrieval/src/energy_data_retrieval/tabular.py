"""Bounded, explicit imports for preserved source exports, without guessing schemas."""

from __future__ import annotations

import csv
import io
import zipfile
from contextlib import ExitStack
from datetime import UTC, datetime, timedelta
from pathlib import Path, PurePosixPath
from zoneinfo import ZoneInfo

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from energy_data_retrieval.transport import RetrievalError

MAX_EXPORT_BYTES = 512 * 1024 * 1024
MAX_EXPORT_ROWS = 2_000_000


def read_export(
    path: Path,
    *,
    member: str | None = None,
    sheet: str | None = None,
    skiprows: int = 0,
    delimiter: str = ",",
    maximum_bytes: int = MAX_EXPORT_BYTES,
    maximum_rows: int = MAX_EXPORT_ROWS,
) -> pd.DataFrame:
    """Read a selected CSV/XLSX table; ZIP members are never extracted to disk.

    Limits apply to source bytes, the selected decompressed member and table rows.
    XLSX cells use cached values with read-only iteration, never evaluated formulas.
    Column names, delimiter, sheet and preamble length are explicit source contracts.
    """
    if maximum_bytes < 1 or maximum_rows < 1 or skiprows < 0 or len(delimiter) != 1:
        raise ValueError("Invalid export resource limits or CSV delimiter")
    if path.stat().st_size > maximum_bytes:
        raise RetrievalError("Export exceeds the configured byte limit")
    suffix = path.suffix.lower()
    if suffix == ".zip":
        if member is None:
            raise ValueError("ZIP input requires an explicit member name")
        name = PurePosixPath(member)
        if name.is_absolute() or ".." in name.parts or "\\" in member:
            raise ValueError("ZIP member must be a safe relative file name")
        with zipfile.ZipFile(path) as archive:
            matches = [entry for entry in archive.infolist() if entry.filename == member]
            if len(matches) != 1 or matches[0].is_dir() or matches[0].flag_bits & 1:
                raise ValueError("ZIP requires exactly one unencrypted selected member")
            if matches[0].file_size > maximum_bytes:
                raise RetrievalError("ZIP member exceeds the configured byte limit")
            with archive.open(matches[0]) as stream:
                payload = stream.read(maximum_bytes + 1)
            if len(payload) > maximum_bytes:
                raise RetrievalError("ZIP member exceeds the configured byte limit")
        source: Path | io.BytesIO = io.BytesIO(payload)
        suffix = name.suffix.lower()
    else:
        if member is not None:
            raise ValueError("--zip-member only applies to ZIP inputs")
        source = path
    if suffix == ".csv":
        if sheet is not None:
            raise ValueError("Sheet selection only applies to XLSX inputs")
        # pandas otherwise silently renames repeated headers, hiding ambiguity.
        with ExitStack() as stack:
            raw_stream = (
                stack.enter_context(source.open("rb")) if isinstance(source, Path) else source
            )
            text_stream = io.TextIOWrapper(raw_stream, encoding="utf-8-sig", newline="")
            stack.callback(text_stream.detach)
            for _ in range(skiprows):
                next(text_stream, None)
            columns = next(csv.reader(text_stream, delimiter=delimiter), None)
            if not columns or any(not name.strip() for name in columns):
                raise ValueError("CSV header requires explicit nonempty column names")
            if len(set(columns)) != len(columns):
                raise ValueError("Export column names must be unique")
        if isinstance(source, io.BytesIO):
            source.seek(0)
        frame = pd.read_csv(
            source,
            sep=delimiter,
            skiprows=skiprows,
            dtype=str,
            keep_default_na=False,
            encoding="utf-8-sig",
            nrows=maximum_rows + 1,
        )
    elif suffix == ".xlsx":
        if sheet is None:
            raise ValueError("XLSX input requires an explicit sheet name")
        # XLSX itself is a ZIP; bound its expanded size before invoking XML readers.
        with zipfile.ZipFile(source) as workbook_zip:
            if sum(item.file_size for item in workbook_zip.infolist()) > maximum_bytes:
                raise RetrievalError("Expanded XLSX exceeds the configured byte limit")
        if isinstance(source, io.BytesIO):
            source.seek(0)
        workbook = load_workbook(source, read_only=True, data_only=True, keep_links=False)
        try:
            if sheet not in workbook.sheetnames:
                raise ValueError("Requested sheet does not exist in the workbook")
            rows = workbook[sheet].iter_rows(values_only=True)
            for _ in range(skiprows):
                next(rows, None)
            header = next(rows, None)
            if header is None or any(item is None for item in header):
                raise ValueError("Workbook header must contain explicit column names")
            names = [str(item) for item in header]
            if len(set(names)) != len(names):
                raise ValueError("Export column names must be unique")
            records: list[list[object]] = []
            for row in rows:
                if len(records) >= maximum_rows:
                    raise RetrievalError("Export exceeds the configured row limit")
                records.append(list(row))
            frame = pd.DataFrame(records, columns=names)
        finally:
            workbook.close()
    else:
        raise ValueError("Supported exports are CSV, XLSX, or a selected CSV/XLSX ZIP member")
    if len(frame) > maximum_rows:
        raise RetrievalError("Export exceeds the configured row limit")
    if not frame.columns.is_unique:
        raise ValueError("Export column names must be unique")
    return frame


def numeric_values(values: pd.Series, *, decimal: str = ".") -> pd.Series:
    """Keep blank/explicit missing values; reject malformed numbers and infinities."""
    if decimal not in {".", ","}:
        raise ValueError("Decimal separator must be '.' or ','")
    tokens = values.astype("string").str.strip()
    tokens = tokens.mask(tokens.isin(["", "NA", "N/A", "null", "NaN", "nan"]))
    if decimal == ",":
        tokens = tokens.str.replace(",", ".", regex=False)
    result = pd.to_numeric(tokens, errors="coerce").astype("float64")
    if (tokens.notna() & result.isna()).any() or np.isinf(result.to_numpy()).any():
        raise ValueError("Source contains malformed or infinite numeric observations")
    return result


def utc_intervals(values: pd.Series, *, timezone: str, label: str) -> pd.Series:
    """Resolve one-hour timestamps without guessing DST folds or nonexistent times."""
    if label not in {"start", "end"}:
        raise ValueError("Timestamp label must be 'start' or 'end'")
    zone = ZoneInfo(timezone)
    cache: dict[str, datetime] = {}
    for raw in values.unique():
        text = str(raw)
        try:
            stamp = datetime.fromisoformat(text)
        except ValueError:
            raise ValueError("Source timestamps must be explicit ISO dates/times") from None
        if stamp.minute or stamp.second or stamp.microsecond:
            raise ValueError("Hourly export labels must lie on hour boundaries")
        if stamp.tzinfo is None:
            candidates = {
                aware.astimezone(UTC)
                for fold in (0, 1)
                if (aware := stamp.replace(tzinfo=zone, fold=fold))
                .astimezone(UTC)
                .astimezone(zone)
                .replace(tzinfo=None)
                == stamp
            }
            if len(candidates) != 1:
                raise ValueError(
                    "Ambiguous or nonexistent civil-clock label; obtain explicit UTC offsets"
                )
            resolved = next(iter(candidates))
        else:
            resolved = stamp.astimezone(UTC)
        cache[text] = resolved - timedelta(hours=int(label == "end"))
    return pd.to_datetime(values.astype(str).map(cache), utc=True)


def identity_mapping(path: Path, columns: tuple[str, ...]) -> pd.DataFrame:
    """Load exact source identifiers without fuzzy name joins or numeric coercion."""
    frame = pd.read_csv(path, dtype=str, keep_default_na=False)
    if not set(columns).issubset(frame.columns) or frame.empty:
        raise ValueError(f"Mapping requires nonempty columns: {', '.join(columns)}")
    frame = frame[list(columns)]
    if frame.isna().any().any() or frame.map(lambda value: not value.strip()).any().any():
        raise ValueError("Mapping identities cannot be missing or blank")
    if frame.duplicated().any():
        raise ValueError("Mapping contains duplicate records")
    return frame
