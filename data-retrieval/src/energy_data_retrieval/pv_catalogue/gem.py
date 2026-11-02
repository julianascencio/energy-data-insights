"""Offline GEM project import: group phases, retain evidence and screen overlaps.

Only explicit PV and unrestricted complete locations are candidates. Coordinates
and dated statuses remain source claims. No network, fuzzy merge or AC conversion.
"""

from __future__ import annotations

import hashlib
import io
import json
import math
import re
import unicodedata
from collections import defaultdict
from datetime import date, datetime
from typing import Any

from openpyxl import load_workbook
from scipy.spatial import KDTree

from .catalogue import canonical_json, stable_id

SHEET = "Utility-Scale (1 MW+)"
HEADERS = {
    "Country/Area",
    "Project Name",
    "Phase Name",
    "Capacity (MW)",
    "Capacity Rating",
    "Technology Type",
    "Status",
    "GEM location ID",
    "GEM phase ID",
    "Other IDs (unit/phase)",
    "Other IDs (location)",
    "Latitude",
    "Longitude",
    "Location accuracy",
    "Wiki URL",
    "Date Last Researched",
}


def read_gem_workbook(content: bytes, *, sha256: str) -> list[dict[str, Any]]:
    """Check exact source bytes and headers before reading utility-scale phases."""
    if hashlib.sha256(content).hexdigest() != sha256:
        raise ValueError("GEM source checksum mismatch")
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        iterator = workbook[SHEET].iter_rows(values_only=True)
        headers = next(iterator)
        if not set(headers) >= HEADERS or len(headers) != len(set(headers)):
            raise ValueError("GEM columns changed; review mapping")
        return [
            {
                key: value.isoformat() if isinstance(value, (date, datetime)) else value
                for key, value in zip(headers, row, strict=True)
            }
            for row in iterator
            if any(row)
        ]
    finally:
        workbook.close()


def _text(value: Any) -> str:
    return (
        "" if value is None or str(value).strip() in {"", "--", "unknown"} else str(value).strip()
    )


def _name(value: Any) -> str:
    text = unicodedata.normalize("NFKD", _text(value).casefold())
    text = "".join(c for c in text if not unicodedata.combining(c))
    text = re.sub(r"\b(solar|photovoltaic|pv|power|farm|plant|park|project|station)\b", " ", text)
    return re.sub(r"[^\w]", "", text)


def _names(row: dict[str, Any], *, native: bool = False) -> set[str]:
    values = (
        [row.get("Project Name"), row.get("Project Name in Local Language / Script")]
        if native
        else [row.get("name"), *row.get("aliases", [])]
    )
    if native:
        values += re.split(r"[;|]", _text(row.get("Other Name(s)")))
    return {name for v in values if (name := _name(v))}


def _xyz(lat: Any, lon: Any) -> tuple[float, float, float] | None:
    if not _text(lat) or not _text(lon):
        return None
    try:
        latitude, longitude = float(lat), float(lon)
    except (ValueError, TypeError):
        return None
    if not math.isfinite(latitude + longitude) or abs(latitude) > 90 or abs(longitude) > 180:
        return None
    a, b = math.radians(latitude), math.radians(longitude)
    return math.cos(a) * math.cos(b), math.cos(a) * math.sin(b), math.sin(a)


def stage_gem_projects(
    phases: list[dict[str, Any]],
    *,
    existing: list[dict[str, Any]],
    country_codes: dict[str, str],
    source_id: str,
    now: str,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    """Return candidates and ID-only hold reasons; never mutate existing records.

    Hold complete locations on mixed technology/licence, ambiguous names, nearby
    existing projects (2 km), or missing identity. All US sites need an EIA crosswalk.
    Different phase statuses/bases are retained without a misleading scalar total.
    Kosovo uses explicitly documented user-assigned XKX, not an ISO assignment.
    """
    groups: dict[str, list[dict[str, Any]]] = defaultdict(list)
    seen = set()
    for row in phases:
        location, phase = _text(row.get("GEM location ID")), _text(row.get("GEM phase ID"))
        if not re.fullmatch(r"L1\d{11}", location) or not re.fullmatch(r"G1\d{11}", phase):
            raise ValueError("Missing or malformed GEM identity")
        if phase in seen:
            raise ValueError("Duplicate GEM phase ID")
        seen.add(phase)
        groups[location].append(row)
    name_groups: dict[tuple[str, str], set[str]] = defaultdict(set)
    for location, group in groups.items():
        for row in group:
            for name in _names(row, native=True):
                name_groups[(_text(row.get("Country/Area")), name)].add(location)
    existing_names: dict[str, set[str]] = defaultdict(set)
    points: dict[str, list[tuple[float, float, float]]] = defaultdict(list)
    existing_gem = set()
    for row in existing:
        if row["entity_type"] != "projects":
            continue
        country = row["country_code"]
        existing_names[country].update(_names(row))
        if point := _xyz(row.get("latitude"), row.get("longitude")):
            points[country].append(point)
        external = json.loads(row.get("external_ids_json", "{}"))
        if value := external.get("gem_location_id"):
            existing_gem.add(value)
    trees = {country: KDTree(values) for country, values in points.items()}
    accepted, held = [], []
    for location, group in sorted(groups.items()):
        reasons = set()
        countries = {_text(r.get("Country/Area")) for r in group}
        country_name = sorted(countries)[0]
        country = country_codes.get(country_name, "")
        names = set().union(*(_names(r, native=True) for r in group))
        if len(countries) != 1 or not re.fullmatch(r"[A-Z]{3}", country):
            reasons.add("country_identity")
        if any(r.get("Technology Type") != "PV" for r in group):
            reasons.add("not_all_explicit_pv")
        if any(
            re.search(
                r"\b(?:TZ|WKSL)(?:\s+ID)?\s*:",
                str(r.get("Other IDs (unit/phase)")) + "," + str(r.get("Other IDs (location)")),
                re.I,
            )
            for r in group
        ):
            reasons.add("restricted_source_lineage")
        if country == "USA":
            reasons.add("us_eia_crosswalk_required")
        if location in existing_gem or names & existing_names[country]:
            reasons.add("existing_identity")
        if not names or any(len(name_groups[(country_name, name)]) > 1 for name in names):
            reasons.add("ambiguous_source_name")
        tree = trees.get(country)
        if tree is not None and any(
            tree.query(point)[0] <= 2 * math.sin(2 / (2 * 6371))
            for r in group
            if (point := _xyz(r.get("Latitude"), r.get("Longitude")))
        ):
            reasons.add("near_existing_project_2km")
        for r in group:
            value = r.get("Capacity (MW)")
            if value is not None and (
                not isinstance(value, (int, float)) or not math.isfinite(value) or value < 0
            ):
                reasons.add("invalid_capacity")
        if reasons:
            held.append(
                {"gem_location_id": location, "phase_count": len(group), "reasons": sorted(reasons)}
            )
            continue
        first = group[0]
        statuses = sorted({_text(r.get("Status")) for r in group})
        row: dict[str, Any] = {
            "record_id": stable_id("projects", "gem_location", location),
            "entity_type": "projects",
            "name": first["Project Name"],
            "country_code": country,
            "technology": "photovoltaic",
            "status": statuses[0] if len(statuses) == 1 else "mixed_phase_status",
            "created_at": now,
            "updated_at": now,
            "last_verified_at": now,
            "sources": [source_id],
            "field_sources_json": canonical_json({"*": source_id}),
            "quality_status": "source_reported",
            "source_release": "February 2026",
            "external_ids_json": canonical_json({"gem_location_id": location}),
            "phases_json": canonical_json(
                [{k: v for k, v in phase.items() if v is not None} for phase in group]
            ),
            "notes": "GEM February 2026 source snapshot; not live plant verification. "
            "One GEM location per row; native phase evidence retained. "
            "Scalar capacity only for complete, uniform status and rating groups.",
        }
        for native, canonical in [
            ("State/Province", "region"),
            ("City", "locality"),
            ("Wiki URL", "product_url"),
        ]:
            values = {_text(r.get(native)) for r in group}
            if len(values) == 1 and (value := next(iter(values))):
                row[canonical] = value
        coordinates = {
            (str(r.get("Latitude")), str(r.get("Longitude")), _text(r.get("Location accuracy")))
            for r in group
        }
        if (
            len(coordinates) == 1
            and _xyz(first.get("Latitude"), first.get("Longitude"))
            and (accuracy := _text(first.get("Location accuracy")))
        ):
            row.update(
                latitude=float(first["Latitude"]),
                longitude=float(first["Longitude"]),
                location_accuracy=f"GEM {accuracy}; assumed WGS84",
            )
        ratings = {r.get("Capacity Rating") for r in group}
        if (
            len(ratings) == 1
            and len(statuses) == 1
            and all(isinstance(r.get("Capacity (MW)"), (int, float)) for r in group)
        ):
            rating = next(iter(ratings))
            key = {
                "MWac": "capacity_ac_mw",
                "MWp/dc": "capacity_dc_mw",
                "unknown": "capacity_unspecified_mw",
            }.get(rating)
            if key:
                row[key] = sum(r["Capacity (MW)"] for r in group)
                row["capacity_basis"] = str(rating)
        accepted.append(row)
    return accepted, held
