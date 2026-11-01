"""Synthetic fixtures verify FIT semantics, identity boundaries and privacy."""

import hashlib
import io
import json
from datetime import datetime

import pytest
from openpyxl import Workbook

from energy_data_retrieval.pv_catalogue.japan_fit import (
    MAIN_HEADERS,
    PARCEL_HEADERS,
    stage_fit_review,
)


def workbook(*, change=None, duplicate=False, parcel_id=None, second=False):
    book = Workbook()
    main = book.active
    main.title = "認定設備"
    for col, label in MAIN_HEADERS.items():
        main.cell(3, col + 1, label)
    row = [
        None,
        "A123456F31",
        "PRIVATE PERSON",
        "PRIVATE REPRESENTATIVE",
        "PRIVATE HOME",
        "PRIVATE PHONE",
        "太陽光",
        1500,
        "PRIVATE PLANT PARCEL",
        0,
        2100,
        datetime(2014, 3, 1),
        datetime(2027, 1, 1),
        "2015年6月",
    ]
    if change:
        row[change[0]] = change[1]
    for col, value in enumerate(row, 1):
        main.cell(5, col, value)
    if duplicate:
        main.append(row)
    if second:
        main.append([*row[:1], "A123457F31", *row[2:]])
    parcels = book.create_sheet("すべての設備所在地")
    for col, label in PARCEL_HEADERS.items():
        parcels.cell(2, col + 1, label)
    parcels.append([None, parcel_id or row[1], "PRIVATE PERSON", 1, "REP", row[6], row[7], row[8]])
    if second:
        parcels.append([None, "A123457F31", "PRIVATE PERSON", 1, "REP", row[6], row[7], row[8]])
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()


def stage(content, **kwargs):
    return stage_fit_review(
        content,
        sha256=hashlib.sha256(content).hexdigest(),
        source_id="synthetic",
        snapshot_date="2026-09-30",
        prefecture="鳥取県",
        **kwargs,
    )


def test_capacity_bases_month_precision_and_forecast_remain_distinct():
    result = stage(workbook())
    item = result["candidates"][0]
    assert result["accepted_count"] == 0
    assert item["registered_generation_output_kw_basis_unspecified"] == 1500
    assert item["solar_cell_total_output_kw"] == 2100
    assert item["reported_operation_month"] == "2015-06"
    assert item["planned_operation_date"] == "2027-01-01"
    assert item["new_certification_date"] == "2014-03-01"
    assert "commissioning_date" not in item and "status" not in item
    assert item["source_locator"] == "認定設備!B5:N5"


def test_private_names_contacts_and_addresses_never_exported():
    result = stage(workbook())
    assert "PRIVATE" not in json.dumps(result)
    assert "REP" not in json.dumps(result)


def test_shared_parcels_flag_ids_without_automatic_merge():
    result = stage(workbook(second=True))
    assert result["candidate_count"] == 2
    assert result["accepted_count"] == 0
    assert result["candidates"][0]["related_registration_ids"] == ["A123457F31"]


@pytest.mark.parametrize("technology", ["バイオマス", "風力", "Solar"])
def test_other_technologies_not_promoted(technology):
    assert stage(workbook(change=(6, technology)))["candidate_count"] == 0


@pytest.mark.parametrize("value", [None, "-", 0])
def test_missing_positive_cell_output_held(value):
    item = stage(workbook(change=(10, value)))["candidates"][0]
    assert not item["photovoltaic_evidence"]
    assert "positive_solar_cell_output_not_verified" in item["review_flags"]


@pytest.mark.parametrize("change", [(7, -1), (10, -1), (13, "2027年1月"), (13, "2015年13月")])
def test_invalid_values_fail_closed(change):
    with pytest.raises(ValueError):
        stage(workbook(change=change))


def test_duplicate_registration_and_unmatched_parcel_fail_closed():
    with pytest.raises(ValueError, match="Duplicate"):
        stage(workbook(duplicate=True))
    with pytest.raises(ValueError, match="mismatch"):
        stage(workbook(parcel_id="A999999F31"))


def test_pinned_bytes_schema_and_scope_are_required():
    content = workbook()
    with pytest.raises(ValueError):
        stage_fit_review(
            content, sha256="0" * 64, source_id="s", snapshot_date="2026-09-30", prefecture="p"
        )
    with pytest.raises(ValueError, match="scope"):
        stage(content, minimum_kw=50)
    book = Workbook()
    stream = io.BytesIO()
    book.save(stream)
    with pytest.raises(ValueError, match="sheets"):
        stage(stream.getvalue())


def test_small_facilities_not_part_of_review_scope():
    result = stage(workbook(change=(7, 999)))
    assert result["facility_rows"] == result["solar_rows"] == 1
    assert result["candidate_count"] == 0
